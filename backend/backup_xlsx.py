"""Back up the whole tournament to an Excel workbook (plus a raw JSON copy).

Reads through the same store the API uses (Neon when DATABASE_URL is set),
so it backs up exactly what the site shows. Read-only - nothing is changed.

    .venv\\Scripts\\python.exe backup_xlsx.py [output_folder]

Easier: run ..\\backup.ps1 from the project folder.
"""

from __future__ import annotations

import json
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app import config, domain
from app.store import get_store

DEFAULT_DIR = config.PROJECT_DIR / "backups"
HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _players(team: dict[str, Any] | None) -> str:
    if not team:
        return ""
    return f"{team['player1_name']} & {team['player2_name']}"


def _super_code(team: dict[str, Any]) -> str | None:
    code, position = team.get("super_group_code"), team.get("super_position")
    return f"{code}{position}" if code and position else None


def _add_sheet(wb: Workbook, title: str, headers: list[str], rows: list[list[Any]]) -> None:
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    for index, header in enumerate(headers, start=1):
        values = [header] + [row[index - 1] for row in rows]
        width = max(len(str(v)) if v is not None else 0 for v in values)
        ws.column_dimensions[get_column_letter(index)].width = min(max(width + 2, 8), 50)


def build_workbook(state: dict[str, Any], exported_at: str, store_name: str) -> Workbook:
    teams = state["teams"]
    matches = state["matches"]
    by_id = {team["id"]: team for team in teams}

    def code_of(team_id: str | None, stage: str) -> str | None:
        team = by_id.get(team_id) if team_id else None
        if not team:
            return None
        return team.get("code") if stage == domain.STAGE_LEAGUE else _super_code(team)

    wb = Workbook()
    wb.remove(wb.active)

    draw = state.get("draw") or {}
    played = sum(1 for m in matches if m.get("played"))
    _add_sheet(
        wb,
        "Summary",
        ["Field", "Value"],
        [
            ["Tournament", config.TOURNAMENT_SLUG],
            ["Backup taken", exported_at],
            ["Data source", store_name],
            ["Current stage", state.get("stage")],
            ["Group draw seed", draw.get("seed")],
            ["Group draw made", draw.get("created_at")],
            ["Super Six seed", state.get("super_six_seed")],
            ["Teams", len(teams)],
            ["Matches", len(matches)],
            ["Matches played", played],
        ],
    )

    _add_sheet(
        wb,
        "Teams",
        [
            "Entry No", "Code", "Player 1", "Player 2", "Group", "Position",
            "Super Six Code", "Super Six Group", "Super Six Position", "Super Six Origin",
            "Team ID",
        ],
        [
            [
                t["entry_no"], t.get("code"), t["player1_name"], t["player2_name"],
                t.get("group_code"), t.get("position"),
                _super_code(t), t.get("super_group_code"), t.get("super_position"),
                t.get("super_origin"), t["id"],
            ]
            for t in teams
        ],
    )

    match_rows = []
    for m in matches:
        stage = m.get("stage") or domain.STAGE_LEAGUE
        winner = m.get("winner_team_id")
        match_rows.append(
            [
                m["match_no"], stage, m.get("slot"), m.get("group_code"), m.get("round"),
                code_of(m.get("team_a_id"), stage), _players(by_id.get(m.get("team_a_id"))),
                code_of(m.get("team_b_id"), stage), _players(by_id.get(m.get("team_b_id"))),
                "Yes" if m.get("played") else "No",
                code_of(winner, stage), _players(by_id.get(winner)),
                m.get("coins"), "Yes" if m.get("red_to_winner") else "No",
                domain.match_pd(m) if m.get("played") else None,
                m["id"], m.get("team_a_id"), m.get("team_b_id"), winner,
            ]
        )
    _add_sheet(
        wb,
        "Matches",
        [
            "Match No", "Stage", "Slot", "Group", "Round",
            "Team A", "Team A Players", "Team B", "Team B Players",
            "Played", "Winner", "Winner Players", "Coins Left", "Red to Winner", "PD",
            "Match ID", "Team A ID", "Team B ID", "Winner ID",
        ],
        match_rows,
    )

    standing_rows = []
    for stage in (domain.STAGE_LEAGUE, domain.STAGE_SUPER_SIX):
        for group in domain.build_groups(teams, matches, stage):
            for row in group["standings"]:
                standing_rows.append(
                    [
                        stage, group["code"], group["name"], row["rank"], row["code"],
                        " & ".join(row["players"]), row["played"], row["won"], row["lost"],
                        row["points"], row["pd"], row["reds"],
                    ]
                )
    _add_sheet(
        wb,
        "Standings",
        ["Stage", "Group", "Group Name", "Rank", "Code", "Players",
         "Played", "Won", "Lost", "Points", "PD", "Reds"],
        standing_rows,
    )
    return wb


def main() -> None:
    out_dir = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_DIR
    out_dir.mkdir(parents=True, exist_ok=True)

    store = get_store()
    try:
        state = store.load()
    finally:
        if hasattr(store, "pool"):
            store.pool.close()
    now = datetime.now()
    exported_at = now.strftime("%Y-%m-%d %H:%M:%S")
    stamp = now.strftime("%Y-%m-%d_%H-%M-%S")

    xlsx_path = out_dir / f"nc-carrom-backup_{stamp}.xlsx"
    json_path = out_dir / f"nc-carrom-backup_{stamp}.json"

    build_workbook(state, exported_at, store.name).save(xlsx_path)
    json_path.write_text(json.dumps(state, indent=2, ensure_ascii=False, default=str), encoding="utf-8")

    print(f"  Teams: {len(state['teams'])}   Matches: {len(state['matches'])}   Stage: {state.get('stage')}")
    print(f"  Saved {xlsx_path}")
    print(f"  Saved {json_path}")


if __name__ == "__main__":
    main()
